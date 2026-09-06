"""Anonymous RT annual archives from ERCOT public report 13061.

Source: https://www.ercot.com/mp/data-products/data-product-details?id=np6-785-er
The seven-column schema was observed in 2010 and 2026 workbooks. Other schema
epochs fail closed; matching headers do not establish complete annual coverage.
"""

import datetime
import hashlib
import io
import json
import os
import re
import zipfile
from collections.abc import Generator, Sequence
from dataclasses import asdict, dataclass
from decimal import Decimal
from pathlib import Path
from typing import Self

import httpx
from pydantic import BaseModel, ConfigDict, Field, StrictInt, StrictStr

from ._http import (
    _HTTP,
    AccessDeniedError,
    LimitError,
    Limits,
    Receipt,
    SchemaMismatchError,
)
from .archives import ArchiveDocument, Download, checked_zip

_LIST_URL = "https://www.ercot.com/misapp/servlets/IceDocListJsonWS"
_DOWNLOAD_URL = "https://www.ercot.com/misdownload/servlets/mirDownload"
_HEADER = (
    "Delivery Date",
    "Delivery Hour",
    "Delivery Interval",
    "Repeated Hour Flag",
    "Settlement Point Name",
    "Settlement Point Type",
    "Settlement Point Price",
)


class RTArchivePrice(BaseModel):
    """Observed public RT workbook fields without an inferred UTC interval.

    Attributes:
        delivery_date: Local source delivery date.
        delivery_hour: Source delivery hour, from 1 through 24.
        delivery_interval: Source quarter-hour interval, from 1 through 4.
        repeated_hour_flag: Original flag, without inferred timezone semantics.
        settlement_point_name: Original public hub or load-zone identifier.
        settlement_point_type: Original source point type.
        settlement_point_price: Decimal conversion of the numeric workbook cell.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)
    delivery_date: datetime.date
    delivery_hour: StrictInt = Field(ge=1, le=24)
    delivery_interval: StrictInt = Field(ge=1, le=4)
    repeated_hour_flag: StrictStr
    settlement_point_name: StrictStr
    settlement_point_type: StrictStr
    settlement_point_price: Decimal = Field(allow_inf_nan=False)


@dataclass(frozen=True)
class RTArchiveRecord:
    """A typed RT price and its physical source location.

    Attributes:
        row: Source price fields.
        sheet: Source worksheet name.
        row_number: One-based physical row including the header.
        member: XLSX member name in the downloaded ZIP.
        source_sha256: Exact downloaded ZIP hash.
    """

    row: RTArchivePrice
    sheet: str
    row_number: int
    member: str
    source_sha256: str


class RTArchiveClient:
    """Anonymous, bounded listing and downloads for public RT report 13061."""

    def __init__(
        self,
        *,
        limits: Limits | None = None,
        transport: httpx.BaseTransport | None = None,
    ) -> None:
        """Create an adapter without making a request.

        Args:
            limits: HTTP budgets; defaults allow a 16 MB annual ZIP.
            transport: Optional mocked transport for offline verification.
        """
        self._http = _HTTP(limits or Limits(max_bytes=16_000_000), transport)
        self._listed: dict[str, ArchiveDocument] = {}

    def __enter__(self) -> Self:
        """Enter a scope that releases connections on exit.

        Returns:
            This client instance.
        """
        return self

    def __exit__(self, *exc: object) -> None:
        """Release connections without suppressing a caller's exception.

        Args:
            *exc: Exception details supplied by the context manager protocol.
        """
        self.close()

    def close(self) -> None:
        """Release connections and prevent future requests."""
        self._http.close()

    def archives(self) -> tuple[tuple[ArchiveDocument, ...], Receipt]:
        """List public annual files, preserving source order and publication time.

        Returns:
            Public documents and a listing receipt, without downloading history.

        Raises:
            PublicDataError: The source is nonpublic, unavailable, or malformed.
        """
        self._listed = {}
        payload = self._http.request("GET", _LIST_URL, params={"reportTypeId": 13061})
        try:
            entries = payload.json()["ListDocsByRptTypeRes"]["DocumentList"]
            if not isinstance(entries, list) or len(entries) > 100:
                raise ValueError
            documents: dict[str, ArchiveDocument] = {}
            for entry in entries:
                item = entry["Document"]
                if item["SecurityStatus"] != "P":
                    raise AccessDeniedError(
                        "RT archive listing contains nonpublic data"
                    )
                identity = item["DocID"]
                name = item["FriendlyName"]
                if (
                    not isinstance(identity, str)
                    or not re.fullmatch(r"[0-9]{1,24}", identity)
                    or not isinstance(name, str)
                    or not re.fullmatch(r"RTMLZHBSPP_[0-9]{4}", name)
                    or item["Extension"].lower() != "zip"
                    or int(item["ReportTypeID"]) != 13061
                    or identity in documents
                ):
                    raise ValueError
                published = datetime.datetime.fromisoformat(item["PublishDate"])
                size = int(item["ContentSize"])
                if published.tzinfo is None or size < 1:
                    raise ValueError
                documents[identity] = ArchiveDocument(
                    identity, 13061, name, published, size, "P"
                )
            self._listed = documents
            return tuple(documents.values()), payload.receipt
        except (KeyError, TypeError, ValueError, AttributeError):
            pass
        raise SchemaMismatchError("RT annual listing differs from the public contract")

    def download(self, document: ArchiveDocument, *, cache: Path) -> Download:
        """Download one listed ZIP or reuse its unchanged, hash-verified receipt.

        Args:
            document: An unchanged entry from this instance's public listing.
            cache: Directory for exclusive, identity-based file and receipt writes.

        Returns:
            One complete ZIP with original publication and retrieval evidence.

        Raises:
            PublicDataError: Access, transfer, ZIP, or cache validation fails.
            OSError: Local cache storage fails. Partial caches require repair.
        """
        if (
            document.security != "P"
            or document.report_type_id != 13061
            or self._listed.get(document.document_id) != document
        ):
            raise AccessDeniedError("RT download requires this client's public listing")
        if document.byte_count > self._http.limits.max_bytes:
            raise LimitError("Advertised RT archive exceeds the download budget")
        cache = Path(cache)
        if cache.is_symlink():
            raise SchemaMismatchError("RT receipt cache cannot be a symlink")
        cache.mkdir(parents=True, exist_ok=True)
        path = cache / f"13061-{document.document_id}.zip"
        receipt_path = path.with_suffix(".json")
        temporary = path.with_suffix(".part")
        if (
            any(p.is_symlink() for p in (path, receipt_path, temporary))
            or temporary.exists()
        ):
            raise SchemaMismatchError("RT cache requires explicit repair")
        url = str(
            httpx.URL(
                _DOWNLOAD_URL,
                params={"doclookupId": document.document_id, "reportTypeId": 13061},
            )
        )
        identity = asdict(document)
        identity["published_at"] = document.published_at.isoformat()
        if path.exists() or receipt_path.exists():
            try:
                if (
                    path.stat().st_size != document.byte_count
                    or receipt_path.stat().st_size > 16_000
                ):
                    raise ValueError
                raw = path.read_bytes()
                saved = json.loads(receipt_path.read_text())
                record = saved["receipt"]
                receipt = Receipt(
                    record["source_url"],
                    datetime.datetime.fromisoformat(record["retrieved_at"]),
                    record["sha256"],
                    record["byte_count"],
                    record["status"],
                )
                if (
                    saved["document"] != identity
                    or receipt.source_url != url
                    or receipt.retrieved_at.tzinfo is None
                    or receipt.status != 200
                    or receipt.byte_count != len(raw)
                    or receipt.sha256 != hashlib.sha256(raw).hexdigest()
                ):
                    raise ValueError
                return Download(document, path, receipt, True)
            except (OSError, KeyError, TypeError, ValueError):
                pass
            raise SchemaMismatchError(
                "RT cache changed or is incomplete; repair explicitly"
            )
        payload = self._http.request("GET", url)
        if payload.receipt.byte_count != document.byte_count:
            raise SchemaMismatchError("RT download size differs from the listing")
        with checked_zip(
            payload.body, max_members=24, max_expanded_bytes=64_000_000
        ) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) != 1 or not members[0].filename.lower().endswith(".xlsx"):
                raise SchemaMismatchError("RT annual ZIP must contain one XLSX")
        descriptor = os.open(
            temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600
        )
        try:
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(payload.body)
                stream.flush()
                os.fsync(stream.fileno())
            os.link(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
        record = asdict(payload.receipt)
        record["retrieved_at"] = payload.receipt.retrieved_at.isoformat()
        descriptor = os.open(
            receipt_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600
        )
        with os.fdopen(descriptor, "w") as stream:
            json.dump({"document": identity, "receipt": record}, stream, indent=2)
            stream.write("\n")
        return Download(document, path, payload.receipt)


def iter_rt_archive(
    download: Download,
    *,
    sheets: Sequence[str] | None = None,
    max_rows: int | None = 1000,
    max_scan_rows: int | None = 100_000,
    max_expanded_bytes: int = 256_000_000,
    max_members: int = 512,
) -> Generator[RTArchiveRecord, None, None]:
    """Stream public RT workbook rows, rejecting unknown columns and cell types.

    Args:
        download: Receipt-verified public report-13061 ZIP.
        sheets: Exact worksheet names, or None for every sheet in source order.
        max_rows: Yield budget; None permits complete selected-file iteration.
        max_scan_rows: Physical row budget; None scans every selected row.
        max_expanded_bytes: Maximum expanded bytes in each ZIP layer.
        max_members: Maximum members in each ZIP layer.

    Yields:
        Typed records preserving duplicates and physical row identities. Close the
        generator when stopping early. Budget exhaustion raises before completion.

    Raises:
        ValueError: A caller-selected bound or worksheet selection is invalid.
        ImportError: The optional files extra is missing.
        PublicDataError: Access, receipt, schema, or archive budget checks fail.
        OSError: The local file cannot be read.
    """
    for value in (max_rows, max_scan_rows):
        if value is not None and (type(value) is not int or value < 1):
            raise ValueError("Row budgets must be positive or None")
    if any(type(v) is not int or v < 1 for v in (max_expanded_bytes, max_members)):
        raise ValueError("Archive structure budgets must be positive")
    if isinstance(sheets, str) or (
        sheets is not None and (not sheets or len(set(sheets)) != len(sheets))
    ):
        raise ValueError("sheets must contain distinct worksheet names")
    document = download.document
    if document.security != "P" or document.report_type_id != 13061:
        raise AccessDeniedError("Only public RT annual files are supported")
    if (
        download.path.is_symlink()
        or download.path.stat().st_size != download.receipt.byte_count
        or download.receipt.byte_count != document.byte_count
        or not re.fullmatch(r"RTMLZHBSPP_[0-9]{4}", document.friendly_name)
    ):
        raise SchemaMismatchError(
            "RT archive identity or size differs from its receipt"
        )
    raw = download.path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != download.receipt.sha256:
        raise SchemaMismatchError("RT archive differs from its receipt hash")
    from openpyxl import load_workbook

    try:
        with checked_zip(
            raw, max_members=max_members, max_expanded_bytes=max_expanded_bytes
        ) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) != 1 or not members[0].filename.lower().endswith(".xlsx"):
                raise SchemaMismatchError("Unsupported RT annual container")
            member = members[0].filename
            inner = archive.read(members[0])
        with checked_zip(
            inner, max_members=max_members, max_expanded_bytes=max_expanded_bytes
        ):
            pass
        workbook = load_workbook(
            io.BytesIO(inner), read_only=True, data_only=True, keep_links=False
        )
        try:
            selected = list(workbook.sheetnames if sheets is None else sheets)
            if any(name not in workbook.sheetnames for name in selected):
                raise SchemaMismatchError("Requested RT worksheet is absent")
            count = scanned = 0
            for name in selected:
                iterator = workbook[name].iter_rows(values_only=True)
                if tuple(next(iterator, ())) != _HEADER:
                    raise SchemaMismatchError("Unknown RT annual schema epoch")
                for number, cells in enumerate(iterator, start=2):
                    scanned += 1
                    if max_scan_rows is not None and scanned > max_scan_rows:
                        raise LimitError("RT scan budget reached before completion")
                    if all(value is None for value in cells):
                        continue
                    if max_rows is not None and count >= max_rows:
                        raise LimitError("RT row budget reached before completion")
                    if len(cells) != 7:
                        raise SchemaMismatchError("RT annual columns changed")
                    day, hour, interval, repeated, point, kind, price = cells
                    if (
                        not isinstance(day, str)
                        or not re.fullmatch(r"[0-9]{2}/[0-9]{2}/[0-9]{4}", day)
                        or type(price) not in (int, float)
                    ):
                        raise SchemaMismatchError("RT annual cell types changed")
                    date = datetime.date(int(day[6:]), int(day[:2]), int(day[3:5]))
                    if date.year != int(document.friendly_name[-4:]):
                        raise SchemaMismatchError(
                            "RT row date differs from the listed year"
                        )
                    row = RTArchivePrice(
                        delivery_date=date,
                        delivery_hour=hour,
                        delivery_interval=interval,
                        repeated_hour_flag=repeated,
                        settlement_point_name=point,
                        settlement_point_type=kind,
                        settlement_point_price=Decimal(str(price)),
                    )
                    count += 1
                    yield RTArchiveRecord(row, name, number, member, digest)
        finally:
            workbook.close()
    except (zipfile.BadZipFile, KeyError, TypeError, ValueError):
        pass
    else:
        return
    raise SchemaMismatchError("RT workbook differs from the observed contract")
