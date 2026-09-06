"""Sampling and opt-in complete iteration of public DAM annual ZIP/XLSX files."""

import datetime
import hashlib
import io
import re
import stat
import zipfile
from collections.abc import Generator, Sequence
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path, PurePosixPath

from pydantic import BaseModel, ConfigDict, StrictStr

from ._http import AccessDeniedError, LimitError, Receipt, SchemaMismatchError

HEADER = (
    "Delivery Date",
    "Hour Ending",
    "Repeated Hour Flag",
    "Settlement Point",
    "Settlement Point Price",
)


@dataclass(frozen=True)
class ArchiveDocument:
    """One public MIS listing entry, not a guarantee of interval completeness.

    Attributes:
        document_id: Source document identity used for receipts and downloads.
        report_type_id: Fixed public DAM annual archive report type, 13060.
        friendly_name: Source year-file label.
        published_at: Source publication timestamp with its original offset.
        byte_count: Advertised download size, separate from the received size.
        security: Raw MIS security code; only P is supported.
    """

    document_id: str
    report_type_id: int
    friendly_name: str
    published_at: datetime.datetime
    byte_count: int
    security: str


@dataclass(frozen=True)
class Download:
    """One verified local archive file and its source receipt.

    Attributes:
        document: Original listing metadata, including publication time.
        path: Complete file, committed atomically after validation.
        receipt: Original download receipt; cache reuse keeps its retrieval time.
        cache_hit: Whether a prior complete hash-verified download was reused.
    """

    document: ArchiveDocument
    path: Path
    receipt: Receipt
    cache_hit: bool = False


class ArchivePrice(BaseModel):
    """A row from the observed DAM annual workbook format.

    Attributes:
        delivery_date: Source delivery date from the workbook.
        hour_ending: Original hour-ending text, including 24:00.
        repeated_hour_flag: Raw workbook flag, without a guessed timezone mapping.
        settlement_point: Original settlement point identifier.
        settlement_point_price: Decimal conversion of the source numeric cell.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)
    delivery_date: datetime.date
    hour_ending: StrictStr
    repeated_hour_flag: StrictStr
    settlement_point: StrictStr
    settlement_point_price: Decimal


@dataclass(frozen=True)
class ArchiveSample:
    """A bounded worksheet selection, never a complete annual-history claim.

    Attributes:
        rows: Decoded source rows up to the caller's limit.
        member: Original XLSX member name in the outer ZIP.
        sheet: Original worksheet name.
        truncated: Whether another data row exists after the sampling limit.
        source_sha256: Exact outer archive hash.
    """

    rows: tuple[ArchivePrice, ...]
    member: str
    sheet: str
    truncated: bool
    source_sha256: str


@dataclass(frozen=True)
class ArchiveRecord:
    """A streamed annual row with its worksheet and file identity.

    Attributes:
        row: Typed source price fields, without inferred UTC market intervals.
        sheet: Source worksheet name.
        row_number: One-based source worksheet row number, including the header.
        member: XLSX member name within the source ZIP.
        source_sha256: Exact downloaded archive hash.
    """

    row: ArchivePrice
    sheet: str
    row_number: int
    member: str
    source_sha256: str


def checked_zip(
    raw: bytes, *, max_members: int, max_expanded_bytes: int
) -> zipfile.ZipFile:
    """Inspect archive structure before reading any member data.

    Args:
        raw: Complete compressed archive within the download byte budget.
        max_members: Maximum member count, including directory entries.
        max_expanded_bytes: Maximum sum of advertised decompressed member sizes.

    Returns:
        An open ZipFile. The caller must close it.

    Raises:
        SchemaMismatchError: ZIP syntax or member paths are unsafe.
        LimitError: A member or expanded-byte ceiling is exceeded.
    """
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise SchemaMismatchError("Source did not return a ZIP archive") from None
    members = archive.infolist()
    if (
        len(members) > max_members
        or sum(m.file_size for m in members) > max_expanded_bytes
    ):
        archive.close()
        raise LimitError("Archive expansion budget exceeded")
    seen = set()
    for member in members:
        path = PurePosixPath(member.filename)
        mode = member.external_attr >> 16
        if (
            path.is_absolute()
            or ".." in path.parts
            or "\\" in member.filename
            or ":" in member.filename
            or member.filename in seen
            or stat.S_ISLNK(mode)
            or member.flag_bits & 1
        ):
            archive.close()
            raise SchemaMismatchError("Unsafe or duplicate archive member")
        seen.add(member.filename)
    return archive


def sample_dam_archive(
    download: Download,
    *,
    sheet: str,
    max_rows: int = 100,
) -> ArchiveSample:
    """Decode a bounded worksheet sample from a verified local annual file.

    Requires the files extra. The decoder checks both ZIP layers before using
    openpyxl read-only mode. It does not extract files or follow external links.

    Args:
        download: A public report-13060 file with matching receipt hash.
        sheet: Exact worksheet name; no automatic all-sheet traversal occurs.
        max_rows: Maximum decoded data rows, from 1 through 1000.

    Returns:
        Source rows, workbook identities, and an explicit truncation flag.

    Raises:
        ImportError: Install tinyercot[files] for the XLSX decoder.
        PublicDataError: Access, hash, shape, or archive limits fail.
        ValueError: The row sampling bound is invalid.
    """
    if type(max_rows) is not int or not 1 <= max_rows <= 1000:
        raise ValueError("max_rows must be between 1 and 1000")
    if download.document.security != "P" or download.document.report_type_id != 13060:
        raise AccessDeniedError("Only public DAM annual files are supported")
    if download.path.is_symlink() or download.path.stat().st_size > 4_000_000:
        raise LimitError("Local archive size or path is invalid")
    raw = download.path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != download.receipt.sha256:
        raise SchemaMismatchError("Archive does not match its download receipt")
    from openpyxl import load_workbook

    try:
        with checked_zip(raw, max_members=24, max_expanded_bytes=32_000_000) as outer:
            members = [m for m in outer.infolist() if not m.is_dir()]
            if len(members) != 1 or not members[0].filename.lower().endswith(".xlsx"):
                raise SchemaMismatchError("Unsupported annual archive container")
            member = members[0].filename
            workbook_bytes = outer.read(members[0])
        with checked_zip(
            workbook_bytes, max_members=512, max_expanded_bytes=64_000_000
        ):
            pass
        workbook = load_workbook(
            io.BytesIO(workbook_bytes), read_only=True, data_only=True, keep_links=False
        )
        try:
            if sheet not in workbook.sheetnames:
                raise SchemaMismatchError("Requested source worksheet is absent")
            iterator = workbook[sheet].iter_rows(values_only=True)
            header = next(iterator)
            if tuple(header) != HEADER:
                raise SchemaMismatchError("Annual workbook headers changed")
            rows = []
            truncated = False
            for scanned, cells in enumerate(iterator, start=1):
                if scanned > max_rows * 2 + 1:
                    raise LimitError("Worksheet scan budget exceeded")
                if all(x is None for x in cells):
                    continue
                if len(rows) == max_rows:
                    truncated = True
                    break
                if len(cells) != 5:
                    raise SchemaMismatchError("Annual workbook has extra columns")
                day, hour, repeated, point, price = cells[:5]
                if (
                    not isinstance(day, str)
                    or not re.fullmatch(r"\d{2}/\d{2}/\d{4}", day)
                    or type(price) not in (int, float)
                ):
                    raise SchemaMismatchError("Annual workbook cell types changed")
                amount = Decimal(str(price))
                if not amount.is_finite():
                    raise SchemaMismatchError("Annual workbook price is nonfinite")
                rows.append(
                    ArchivePrice(
                        delivery_date=datetime.date(
                            int(day[6:10]), int(day[:2]), int(day[3:5])
                        ),
                        hour_ending=hour,
                        repeated_hour_flag=repeated,
                        settlement_point=point,
                        settlement_point_price=amount,
                    )
                )
            return ArchiveSample(tuple(rows), member, sheet, truncated, digest)
        finally:
            workbook.close()
    except (zipfile.BadZipFile, KeyError, TypeError, ValueError, StopIteration):
        pass
    raise SchemaMismatchError("Annual workbook differs from the observed XLSX contract")


def iter_dam_archive(
    download: Download,
    *,
    sheets: Sequence[str] | None = None,
    max_rows: int | None = 1000,
    max_scan_rows: int | None = 100_000,
    max_expanded_bytes: int = 64_000_000,
    max_members: int = 512,
) -> Generator[ArchiveRecord, None, None]:
    """Stream selected worksheets or a complete public annual workbook.

    Use max_rows=None and max_scan_rows=None for complete row iteration. These
    are caller choices, not a permanent sampling ceiling. The generator closes
    the workbook on exhaustion or close(); close it when stopping early.

    Args:
        download: Public report-13060 ZIP with a matching receipt hash.
        sheets: Exact worksheet names, or None for all worksheets in source order.
        max_rows: Total yielded-row budget, or None for all selected data rows.
        max_scan_rows: Data/blank row scan budget, or None to scan complete sheets.
        max_expanded_bytes: Caller-selected bound on each ZIP layer's expansion.
        max_members: Caller-selected bound on each ZIP layer's member count.

    Yields:
        Typed prices with physical row and worksheet identities. Budget exhaustion
        raises explicitly; it does not silently claim a complete annual result.

    Raises:
        ValueError: A caller-selected bound or sheet selection is invalid.
        ImportError: The optional files extra is not installed.
        PublicDataError: Source access, hash, ZIP, cell, or budget checks fail.
        OSError: The local archive cannot be read.
    """
    for value in (max_rows, max_scan_rows):
        if value is not None and (type(value) is not int or value < 1):
            raise ValueError("Row and scan budgets must be positive or None")
    if any(
        type(value) is not int or value < 1
        for value in (max_expanded_bytes, max_members)
    ):
        raise ValueError("Archive structure budgets must be positive")
    if isinstance(sheets, str) or (
        sheets is not None and (not sheets or len(set(sheets)) != len(sheets))
    ):
        raise ValueError("sheets must contain distinct worksheet names")
    if download.document.security != "P" or download.document.report_type_id != 13060:
        raise AccessDeniedError("Only public DAM annual files are supported")
    if (
        download.path.is_symlink()
        or download.path.stat().st_size != download.receipt.byte_count
    ):
        raise SchemaMismatchError("Archive path or size differs from its receipt")
    raw = download.path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != download.receipt.sha256:
        raise SchemaMismatchError("Archive differs from its receipt hash")
    from openpyxl import load_workbook

    try:
        with checked_zip(
            raw, max_members=max_members, max_expanded_bytes=max_expanded_bytes
        ) as outer:
            members = [member for member in outer.infolist() if not member.is_dir()]
            if len(members) != 1 or not members[0].filename.lower().endswith(".xlsx"):
                raise SchemaMismatchError("Unsupported annual archive container")
            member = members[0].filename
            inner = outer.read(members[0])
        with checked_zip(
            inner, max_members=max_members, max_expanded_bytes=max_expanded_bytes
        ):
            pass
        workbook = load_workbook(
            io.BytesIO(inner), read_only=True, data_only=True, keep_links=False
        )
        try:
            selected = list(workbook.sheetnames if sheets is None else sheets)
            if any(name not in workbook.sheetnames for name in selected):
                raise SchemaMismatchError("Requested source worksheet is absent")
            count = scanned = 0
            for name in selected:
                iterator = workbook[name].iter_rows(values_only=True)
                header = next(iterator, None)
                if header is None:
                    continue
                if tuple(header) != HEADER:
                    raise SchemaMismatchError("Annual workbook headers changed")
                for row_number, cells in enumerate(iterator, start=2):
                    scanned += 1
                    if max_scan_rows is not None and scanned > max_scan_rows:
                        raise LimitError(
                            "Worksheet scan budget reached before completion"
                        )
                    if all(value is None for value in cells):
                        continue
                    if max_rows is not None and count >= max_rows:
                        raise LimitError("Annual row budget reached before completion")
                    if len(cells) != 5:
                        raise SchemaMismatchError("Annual workbook has extra columns")
                    day, hour, repeated, point, price = cells
                    if (
                        not isinstance(day, str)
                        or not re.fullmatch(r"\d{2}/\d{2}/\d{4}", day)
                        or type(price) not in (int, float)
                    ):
                        raise SchemaMismatchError("Annual workbook cell types changed")
                    amount = Decimal(str(price))
                    if not amount.is_finite():
                        raise SchemaMismatchError("Annual workbook price is nonfinite")
                    row = ArchivePrice(
                        delivery_date=datetime.date(
                            int(day[6:10]), int(day[:2]), int(day[3:5])
                        ),
                        hour_ending=hour,
                        repeated_hour_flag=repeated,
                        settlement_point=point,
                        settlement_point_price=amount,
                    )
                    count += 1
                    yield ArchiveRecord(row, name, row_number, member, digest)
        finally:
            workbook.close()
    except (zipfile.BadZipFile, KeyError, TypeError, ValueError):
        pass
    else:
        return
    raise SchemaMismatchError("Annual workbook differs from the observed XLSX contract")
