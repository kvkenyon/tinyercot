"""Monthly forecast errors and backcasts with original date and scale metadata."""

from __future__ import annotations

import re
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Literal, get_args

from ._legacy_load import _number
from ._load import _workbooks
from ._load_forecasts import _ForecastRow, _ForecastTable
from ._public_tables import PublicFile

ForecastPerformanceSeries = Literal[
    "Day-Ahead RUC",
    "6 Hour-Ahead",
    "3 Hour-Ahead",
    "Goal",
    "Stretch",
    "Backcast",
    "Day-Ahead Backcast",
]


class MonthlyLoadForecastPerformance(_ForecastRow):
    """One month/series, keeping raw values and unlabelled scales distinct."""

    month: date
    kind: Literal["forecast", "backcast", "target"]
    series: ForecastPerformanceSeries
    value: Decimal | None
    sourceDate: date | None
    sourceNumberFormat: str
    sourceMarker: str | None
    sourceColumn: int
    sourceMonthColumn: int
    sourceNotes: list[str]

    @property
    def percent(self) -> Decimal | None:
        """Convert explicit Excel percentage fractions; unlabelled scales stay unknown."""
        if self.value is not None and re.fullmatch(
            r"0(?:\.0+)?%", self.sourceNumberFormat
        ):
            return self.value * 100
        return None


class MonthlyForecastPerformance(_ForecastTable[MonthlyLoadForecastPerformance]):
    title_pattern = r".*Backcast.*"
    extensions = (".xlsx",)

    def files(self) -> list[PublicFile]:
        # Jan/Feb 2026 Backcast links instead contain hourly Metrics workbooks.
        return [
            f for f in super().files() if "backcast" in f.url.rsplit("/", 1)[-1].lower()
        ]

    def _read(
        self, data: bytes, filename: str
    ) -> Iterator[MonthlyLoadForecastPerformance]:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ImportError(
                "Install tinyercot[files] to read monthly forecast performance"
            ) from error
        found = False
        for member, content in _workbooks(data):
            member = filename if member == "workbook.xlsx" else member
            book = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                for sheet in book:
                    rows = list(sheet.iter_rows())
                    notes = [
                        str(c.value)
                        for row in rows[:3]
                        for c in row
                        if isinstance(c.value, str)
                        and c.value not in ("Data Source", "`")
                    ]
                    for header_row, header in enumerate(rows):
                        for column, cell in enumerate(header):
                            if cell.value != "Month":
                                continue
                            columns = []
                            for value_column in range(column + 1, len(header)):
                                label = header[value_column].value
                                if label in (None, ""):
                                    break
                                if label not in get_args(ForecastPerformanceSeries):
                                    raise ValueError(
                                        f"{member}/{sheet.title}: unknown monthly series {label!r}"
                                    )
                                columns.append((value_column, str(label)))
                            if not columns:
                                raise ValueError(
                                    f"{member}/{sheet.title}: monthly table has no series"
                                )
                            year = (
                                rows[header_row - 1][column].value
                                if header_row
                                else None
                            )
                            for line in range(header_row + 1, len(rows)):
                                raw_month = rows[line][column].value
                                if raw_month in (None, ""):
                                    break
                                source_date = (
                                    raw_month.date()
                                    if isinstance(raw_month, datetime)
                                    else raw_month
                                    if isinstance(raw_month, date)
                                    else None
                                )
                                if source_date is not None:
                                    month = source_date.replace(day=1)
                                elif (
                                    isinstance(raw_month, int)
                                    and 1 <= raw_month <= 12
                                    and isinstance(year, int)
                                ):
                                    month = date(year, raw_month, 1)
                                else:
                                    raise ValueError(
                                        f"{member}/{sheet.title}:{line + 1}: invalid month or year"
                                    )
                                for value_column, label in columns:
                                    value_cell = rows[line][value_column]
                                    marker = (
                                        str(value_cell.value)
                                        if value_cell.data_type == "e"
                                        else None
                                    )
                                    kind = (
                                        "target"
                                        if label in ("Goal", "Stretch")
                                        else "backcast"
                                        if "Backcast" in label
                                        else "forecast"
                                    )
                                    found = True
                                    yield MonthlyLoadForecastPerformance.model_validate(
                                        {
                                            "month": month,
                                            "kind": kind,
                                            "series": label,
                                            "value": None
                                            if marker
                                            else _number(value_cell.value),
                                            "sourceDate": source_date,
                                            "sourceNumberFormat": value_cell.number_format,
                                            "sourceMarker": marker,
                                            "sourceColumn": value_column + 1,
                                            "sourceMonthColumn": column + 1,
                                            "sourceNotes": notes,
                                            "sourceMember": member,
                                            "sourceSheet": sheet.title,
                                            "sourceRow": line + 1,
                                        }
                                    )
            finally:
                book.close()
        if not found:
            raise ValueError("No monthly forecast performance tables found")
