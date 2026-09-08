"""Typed site keys accompanying public wind and solar planning profiles."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import TYPE_CHECKING, Literal

from pydantic import BaseModel, ConfigDict

if TYPE_CHECKING:
    from openpyxl.cell.cell import Cell, MergedCell
    from openpyxl.worksheet.worksheet import Worksheet


class ProfileKeyArchive(BaseModel):
    """A key workbook for one study and fuel, as linked by ERCOT."""

    studyYear: int
    yearFrom: int
    yearTo: int
    fuel: Literal["wind", "solar"]
    title: str
    url: str


class ProfileSite(BaseModel):
    """A site or unit exactly as described by its key's table.

    IDs and capacities belong to this key version. Solar table sections retain
    their original numbering without inferring development status from position.
    `markedAsQueued` follows the wind key's explicit cell-shading legend; None
    means that no such legend is supplied. Flags retain their source year.
    """

    model_config = ConfigDict(extra="forbid")
    kind: Literal[
        "solar-operational-planned",
        "solar-hypothetical",
        "solar-metro",
        "solar-rural",
        "wind-operational-planned",
        "wind-hypothetical",
    ]
    sourceSheet: str
    sourceRow: int
    sourceSection: int
    siteId: int | None = None
    unitCode: str | None = None
    commonName: str | None = None
    county: str | None = None
    cdrZone: str | None = None
    capacityMW: Decimal | None = None
    latitude: Decimal | None = None
    longitude: Decimal | None = None
    typeCode: int | None = None
    trackingSystem: str | None = None
    trackingType: str | None = None
    tilt: Decimal | Literal["Lat", "NA"] | None = None
    azimuth: Decimal | None = None
    dcAcRatio: Decimal | None = None
    inverters: str | None = None
    modules: str | None = None
    modeledInYear: int | None = None
    modeledInYearFlag: str | None = None
    newForYear: int | None = None
    newForYearFlag: str | None = None
    developmentStatus: str | None = None
    distanceToExistingOrQueued: str | None = None
    iecClass: int | None = None
    metroArea: str | None = None
    developmentIntensity: str | None = None
    markedAsQueued: bool | None = None


class ProfileUnitMapping(BaseModel):
    """One published wind unit-to-site mapping; one site can have multiple units."""

    model_config = ConfigDict(extra="forbid")
    commonName: str
    unitCode: str
    siteId: int
    newForYear: int
    newForYearFlag: str
    markedAsQueued: bool | None
    sourceSheet: str
    sourceRow: int


class ProfileKeySummary(BaseModel):
    """A printed summary row; totals are not recalculated from site records."""

    model_config = ConfigDict(extra="forbid")
    section: str
    typeCode: int | None
    label: str | None
    profileCount: int | None
    trackingScenarioCount: int | None
    trackingTypes: str | None
    totalProfileCount: int | None
    capacityMW: Decimal | None
    sourceRow: int


class ProfileKeyNote(BaseModel):
    model_config = ConfigDict(extra="forbid")
    text: str
    sourceSheet: str
    sourceCell: str


class ProfileKey(BaseModel):
    """Site tables, mappings, printed summaries and modeling notes from one key."""

    model_config = ConfigDict(extra="forbid")
    sourceMember: str
    title: str | None = None
    revisionDate: date | None = None
    authors: str | None = None
    sites: list[ProfileSite]
    units: list[ProfileUnitMapping]
    summaries: list[ProfileKeySummary]
    notes: list[ProfileKeyNote]


_SITE_TABLES = {
    "Utility_Operational&Planned": "solar-operational-planned",
    "Utility_Hypothetical": "solar-hypothetical",
    "DGPV_metro": "solar-metro",
    "DGPV_rural": "solar-rural",
    "Op&Planned_Wind-Summary": "wind-operational-planned",
    "Hypo_Wind-Summary": "wind-hypothetical",
}
_FIELDS = {
    "Unit Code": "unitCode",
    "SITE ID": "siteId",
    "UL ID": "siteId",
    "Common Name": "commonName",
    "County": "county",
    "COUNTY": "county",
    "CDR Zone": "cdrZone",
    "CDR_Zone": "cdrZone",
    "MWAC": "capacityMW",
    "Capacity_MW": "capacityMW",
    "Modeled Capacity MW (2022)": "capacityMW",
    "Latitude": "latitude",
    "YCoord": "latitude",
    "Longitude": "longitude",
    "XCoord": "longitude",
    "Type": "typeCode",
    "Tracking System": "trackingSystem",
    "Tracking Type": "trackingType",
    "Tilt": "tilt",
    "Azimuth": "azimuth",
    "DC:AC Ratio": "dcAcRatio",
    "Inverter(s)": "inverters",
    "Module(s)": "modules",
    "Modeled in 2020": "modeledInYearFlag",
    "New for 2020": "newForYearFlag",
    "New for 2021": "newForYearFlag",
    "New for 2022": "newForYearFlag",
    "Development Status (2022)": "developmentStatus",
    "Distance to Op/Q'ed": "distanceToExistingOrQueued",
    "IEC Class*": "iecClass",
    "Metro Area": "metroArea",
    "Development Intensity": "developmentIntensity",
    "Common Name of Project": "commonName",
    "RARF UNIT_CODE": "unitCode",
}


def read_key(data: bytes, *, filename: str = "workbook") -> ProfileKey:
    """Read a saved XLSX key with tinyercot[files], including its shading legend."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportError("Install tinyercot[files] to read profile keys") from error
    book = load_workbook(BytesIO(data), data_only=True)
    result = ProfileKey(
        sourceMember=filename, sites=[], units=[], summaries=[], notes=[]
    )
    try:
        for sheet in book:
            if sheet.title == "SUMMARY":
                result.title = str(sheet["B1"].value)
                stamp = sheet["B2"].value
                result.revisionDate = (
                    stamp.date() if isinstance(stamp, datetime) else None
                )
                result.authors = str(sheet["C2"].value)
                result.summaries.extend(_summaries(sheet))
            elif (
                sheet.title in _SITE_TABLES
                or sheet.title == "Op&Planned_Wind-Unit_Codes"
            ):
                _records(sheet, result)
            else:
                raise ValueError(
                    f"{filename}: Unknown profile key worksheet {sheet.title!r}"
                )
    finally:
        book.close()
    if not result.sites:
        raise ValueError("No profile site tables found")
    return result


def _records(sheet: Worksheet, result: ProfileKey) -> None:
    header: tuple[Cell | MergedCell, ...] = ()
    section = 0
    legend = next(
        (
            c
            for row in sheet
            for c in row
            if isinstance(c.value, str)
            and c.value.startswith("Modeled as a Queued Plant")
        ),
        None,
    )
    mapping = sheet.title == "Op&Planned_Wind-Unit_Codes"
    for cells in sheet:
        if not any(c.value is not None for c in cells):
            continue
        first = cells[0].value
        if first in ("Unit Code", "SITE ID", "UL ID", "Common Name of Project"):
            # Side notes lie after the contiguous table headings.
            end = next((i for i, c in enumerate(cells) if c.value is None), len(cells))
            header = cells[:end]
            unknown = [c.value for c in header if str(c.value).strip() not in _FIELDS]
            if unknown:
                raise ValueError(f"{sheet.title}: Unknown key columns {unknown}")
            section += 1
            _notes(cells[end:], sheet.title, result)
            continue
        if not header:
            raise ValueError(f"{sheet.title}: Missing profile key header")
        if first is None or all(c.value is None for c in cells[1 : len(header)]):
            _notes(cells, sheet.title, result)
            continue
        values: dict[str, object] = {
            "sourceSheet": sheet.title,
            "sourceRow": cells[0].row,
        }
        for h, c in zip(header, cells, strict=False):
            label = str(h.value).strip()
            values[_FIELDS[label]] = c.value
            if label.startswith("New for "):
                values["newForYear"] = int(label[-4:])
            elif label.startswith("Modeled in "):
                values["modeledInYear"] = int(label[-4:])
            if label == "UL ID":
                values["markedAsQueued"] = (
                    (
                        c.fill.patternType == legend.fill.patternType
                        and c.fill.fgColor == legend.fill.fgColor
                    )
                    if legend is not None
                    else None
                )
        if mapping:
            result.units.append(ProfileUnitMapping.model_validate(values))
        else:
            result.sites.append(
                ProfileSite.model_validate(
                    values
                    | {
                        "kind": _SITE_TABLES[sheet.title],
                        "sourceSection": section,
                    }
                )
            )
        _notes(cells[len(header) :], sheet.title, result)


def _notes(
    cells: tuple[Cell | MergedCell, ...], sheet: str, result: ProfileKey
) -> None:
    for cell in cells:
        if cell.value is not None:
            if not isinstance(cell.value, str):
                raise ValueError(f"{sheet}/{cell.coordinate}: Unexpected key value")
            result.notes.append(
                ProfileKeyNote(
                    text=cell.value, sourceSheet=sheet, sourceCell=cell.coordinate
                )
            )


def _summaries(sheet: Worksheet) -> Iterator[ProfileKeySummary]:
    section = ""
    layout = ""
    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        _, code, label, count, scenarios, tracking, total = row
        if count == "# of Profiles":
            section = str(label)
            layout = (
                "types"
                if code == "Type"
                else "capacity"
                if scenarios == "MWAC"
                else "overview"
            )
            continue
        if not section or not any(isinstance(v, (int, float)) for v in row):
            continue
        yield ProfileKeySummary.model_validate(
            {
                "section": section,
                "typeCode": code if isinstance(code, int) else None,
                "label": label if label is not None else code,
                "profileCount": count,
                "trackingScenarioCount": scenarios if layout == "types" else None,
                "trackingTypes": tracking if layout == "types" else None,
                "totalProfileCount": total,
                "capacityMW": scenarios if layout == "capacity" else None,
                "sourceRow": number,
            }
        )
