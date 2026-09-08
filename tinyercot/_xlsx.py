"""Optional spreadsheet support for generated historical row models."""

from __future__ import annotations

import re
from collections.abc import Iterator
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from time import strptime
from typing import TypeVar

from pydantic import BaseModel

from ._client import Transport
from ._history import Archive, _archive_files

T = TypeVar("T", bound=BaseModel)


class WorkbookArchive(Archive[T]):
    """Typed report worksheets from ERCOT XLSX downloads (requires ``files`` extra)."""

    def __init__(
        self,
        client: Transport,
        product: str,
        row: type[T],
        columns: dict[str, str],
        dates: dict[str, str],
        *,
        sheets: tuple[str, ...],
        variants: tuple[dict[str, str], ...] = (),
    ) -> None:
        super().__init__(client, product, row, columns, dates, variants=variants)
        self._sheets = sheets

    def read(self, data: bytes) -> Iterator[T]:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ImportError(
                "Install tinyercot[files] to read XLSX archives"
            ) from error
        found = False
        for filename, content in _archive_files(data, "*.xlsx"):
            book = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                for sheet in book:
                    if sheet.title not in self._sheets:
                        continue
                    found = True
                    sheet.reset_dimensions()  # Some ERCOT files incorrectly declare A1:A1.
                    header: tuple[str, ...] | None = None
                    for line, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                        values = list(cells)
                        while values and values[-1] is None:
                            values.pop()
                        if not values:
                            continue
                        if header is None:
                            if (
                                all(
                                    isinstance(v, str) and v in self._columns
                                    for v in values
                                )
                                and len(values) > 1
                            ):
                                header = tuple(v for v in values if isinstance(v, str))
                            continue
                        # Report-builder footers: a date, page number and clock;
                        # older reports use a single formatted text line instead.
                        nonempty = [v for v in values if v is not None]
                        if (
                            isinstance(values[0], datetime)
                            and any(isinstance(v, time) for v in nonempty)
                            and all(
                                isinstance(v, (datetime, time, int)) for v in nonempty
                            )
                        ):
                            continue
                        if (
                            len(values) == 1
                            and isinstance(values[0], str)
                            and (
                                values[0].strip() == "ERCOT Confidential"
                                or re.search(r" - \d+ - ", values[0])
                            )
                        ):
                            continue
                        try:
                            if len(values) > len(header):
                                raise ValueError("Unexpected worksheet columns")
                            values += [None] * (len(header) - len(values))
                            converted: dict[str, object] = {"sourceSheet": sheet.title}
                            for column, value in zip(header, values, strict=True):
                                target = self._columns[column]
                                if isinstance(value, str):
                                    value = value.strip() or None
                                if value is not None and target in self._dates:
                                    value = date(
                                        *strptime(str(value), self._dates[target])[:3]
                                    )
                                elif isinstance(value, float):
                                    value = Decimal(str(value))
                                converted[target] = value
                            yield self._row.model_validate(converted)
                        except ValueError as error:
                            raise ValueError(
                                f"{filename}/{sheet.title}:{line}: {error}"
                            ) from error
                    if header is None:
                        raise ValueError(
                            f"{filename}/{sheet.title}: no supported header"
                        )
            finally:
                book.close()
        if not found:
            raise ValueError("Download contains no supported worksheets")
