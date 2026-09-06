"""Prove full worksheet iteration offline without extracting real history."""

import datetime
import hashlib
import io
import zipfile
from collections import Counter
from pathlib import Path

import pytest

from tinyercot.public import (
    ArchiveDocument,
    Download,
    LimitError,
    Receipt,
    iter_dam_archive,
)
from tinyercot.public.archives import HEADER


def archive(tmp_path: Path) -> Download:
    """Build a small synthetic two-worksheet file with more than 1,000 rows.

    Args:
        tmp_path: Pytest temporary directory for the fixture ZIP.

    Returns:
        A synthetic public document with a valid byte/hash receipt.
    """
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    for index, name in enumerate(("Jan", "Feb")):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        sheet.append(HEADER)
        for number in range(1005):
            sheet.append(["01/01/2026", "24:00", "N", "HB_HOUSTON", number])
    inner = io.BytesIO()
    book.save(inner)
    book.close()
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w", compression=zipfile.ZIP_DEFLATED) as stream:
        stream.writestr("annual.xlsx", inner.getvalue())
    raw = outer.getvalue()
    path = tmp_path / "annual.zip"
    path.write_bytes(raw)
    stamp = datetime.datetime(2026, 9, 5, tzinfo=datetime.UTC)
    return Download(
        ArchiveDocument("123", 13060, "DAMLZHBSPP_2026", stamp, len(raw), "P"),
        path,
        Receipt(
            "https://www.ercot.com/public-test",
            stamp,
            hashlib.sha256(raw).hexdigest(),
            len(raw),
        ),
    )


def test_all_worksheets_and_rows_have_no_permanent_sample_ceiling(tmp_path):
    download = archive(tmp_path)
    rows = iter_dam_archive(download, max_rows=None, max_scan_rows=None)
    counts = Counter()
    for record in rows:
        counts[record.sheet] += 1
        assert record.source_sha256 == download.receipt.sha256
        assert record.row_number == counts[record.sheet] + 1
    assert counts == {"Jan": 1005, "Feb": 1005}


@pytest.mark.parametrize("budget", ["rows", "scan"])
def test_iterator_budgets_raise_instead_of_claiming_completion(tmp_path, budget):
    download = archive(tmp_path)
    with pytest.raises(LimitError):
        list(
            iter_dam_archive(
                download,
                max_rows=2 if budget == "rows" else None,
                max_scan_rows=2 if budget == "scan" else None,
            )
        )


def test_early_close_releases_workbook_and_selected_sheet_is_preserved(
    tmp_path, monkeypatch
):
    from openpyxl.workbook.workbook import Workbook

    download = archive(tmp_path)
    closed = []
    original = Workbook.close

    def close(book):
        closed.append(True)
        original(book)

    monkeypatch.setattr(Workbook, "close", close)
    rows = iter_dam_archive(download, sheets=("Feb",), max_rows=None)
    assert next(rows).sheet == "Feb"
    rows.close()
    assert closed == [True]
