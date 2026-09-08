from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import openpyxl
import pytest

from tinyercot import Client, CrrTimeOfUse, PolrUsage

FIXTURES = Path(__file__).resolve().parents[1] / "tools/inputs/public-tables"


def test_all_original_crr_hours_and_dst_months():
    data = (FIXTURES / "crr-hours.xlsx").read_bytes()
    with Client() as client:
        rows = list(client.crr_hours.read(data, filename="calendar.xlsx"))
    book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        source = list(book.active.values)[2:]
        assert len(rows) == len(source) == 48
        for row, cells in zip(rows, source, strict=True):
            assert row.month == date(
                int(cells[0]),
                (
                    "January",
                    "February",
                    "March",
                    "April",
                    "May",
                    "June",
                    "July",
                    "August",
                    "September",
                    "October",
                    "November",
                    "December",
                ).index(cells[1])
                + 1,
                1,
            )
            assert (
                row.offPeakHours,
                row.peakWDHours,
                row.peakWEHours,
                row.totalHours,
            ) == cells[2:]
            assert row.sourceMember == "calendar.xlsx"
            assert row.sourceSheet == book.active.title
    finally:
        book.close()
    by_month = {r.month: r for r in rows}
    assert by_month[date(2026, 3, 1)].totalHours == 743
    assert by_month[date(2026, 11, 1)].totalHours == 721
    assert by_month[date(2028, 2, 1)].totalHours == 696
    assert CrrTimeOfUse.model_validate_json(rows[0].model_dump_json()) == rows[0]


def test_all_original_polr_counts_energy_and_distinct_periods():
    data = (FIXTURES / "polr.xlsx").read_bytes()
    with Client() as client:
        rows = list(client.polr.read(data, filename="retail.xlsx"))
    book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        source = [
            cells
            for cells in book["PUCT_Even_Year"].values
            if isinstance(cells[2], (int, float))
        ]
        notes = [
            c
            for cells in book["PUCT_Even_Year_Cover_Page"].values
            for c in cells
            if c is not None
        ]
        assert len(rows) == len(source) == 24
        for row, cells in zip(rows, source, strict=True):
            assert row.territory == cells[0] and row.premiseType == cells[1]
            assert row.activeEsiids == cells[2]
            assert row.energyKWh == Decimal(str(cells[3]))
            assert row.snapshotDate == date(2026, 3, 31)
            assert row.energyPeriodStart == date(2025, 4, 1)
            assert row.energyPeriodEnd == date(2026, 3, 31)
            assert row.sourceNotes == notes
            assert (
                row.sourceMember == "retail.xlsx"
                and row.sourceSheet == "PUCT_Even_Year"
            )
    finally:
        book.close()
    assert len({r.territory for r in rows}) == 6
    assert len({r.premiseType for r in rows}) == 4
    assert rows[0].activeEsiids == 824562
    assert rows[0].energyKWh == Decimal(10627695505)
    assert PolrUsage.model_validate_json(rows[0].model_dump_json()) == rows[0]


@pytest.mark.parametrize(
    "service,fixture,title,expected",
    [
        ("crr_hours", "crr-hours", "CRR Time of Use Hours", 48),
        ("polr", "polr", "POLR Counts Energy 2026 Report Final", 24),
    ],
)
def test_live_style_discovery_download_and_saved_file_identity(
    service, fixture, title, expected
):
    url = f"https://www.ercot.com/files/docs/2020/01/01/{fixture}-updated.xlsx"
    requests = []

    def handler(request):
        requests.append(str(request.url))
        assert "authorization" not in request.headers
        if str(request.url) == url:
            return httpx.Response(
                200, content=(FIXTURES / f"{fixture}.xlsx").read_bytes()
            )
        return httpx.Response(200, text=f'<a href="{url}"><span>{title}</span></a>')

    with (
        httpx.Client(transport=httpx.MockTransport(handler)) as http,
        Client(client=http) as client,
    ):
        reader = getattr(client, service)
        rows = list(reader.rows())
        assert requests == [reader.index_url, url]
    assert len(rows) == expected
    assert all(r.sourceMember == f"{fixture}-updated.xlsx" for r in rows)


def test_typed_predicates_filter_source_periods():
    with Client() as client:
        months = list(
            client.crr_hours.read(
                (FIXTURES / "crr-hours.xlsx").read_bytes(),
                where=lambda r: date(2027, 1, 1) <= r.month <= date(2027, 12, 1),
            )
        )
        usage = list(
            client.polr.read(
                (FIXTURES / "polr.xlsx").read_bytes(),
                where=lambda r: (
                    r.snapshotDate == date(2026, 3, 31)
                    and r.premiseType == "Residential"
                ),
            )
        )
    assert len(months) == 12
    assert len(usage) == 6


def test_overlapping_archives_keep_separate_members():
    data = BytesIO()
    with ZipFile(data, "w") as z:
        for name in ("original.xlsx", "revised.xlsx"):
            z.writestr(name, (FIXTURES / "polr.xlsx").read_bytes())
    with Client() as client:
        rows = list(client.polr.read(data.getvalue()))
    assert len(rows) == 48
    assert {r.sourceMember for r in rows} == {"original.xlsx", "revised.xlsx"}


def test_polr_periods_come_from_headers_and_notes_survive_sheet_reordering():
    book = openpyxl.load_workbook(FIXTURES / "polr.xlsx")
    sheet = book["PUCT_Even_Year"]
    for row in sheet:
        if row[0].value == "POLR Territory":
            row[2].value = "# Active ESIIDs Total (3/31/2024)"
            row[3].value = "Annual kWh (4/1/2023 - 3/31/2024)"
    book.move_sheet(sheet, offset=-1)
    data = BytesIO()
    book.save(data)
    book.close()
    with Client() as client:
        rows = list(client.polr.read(data.getvalue(), filename="2026.xlsx"))
    assert all(r.snapshotDate == date(2024, 3, 31) for r in rows)
    assert all(r.energyPeriodStart == date(2023, 4, 1) for r in rows)
    assert all(r.energyPeriodEnd == date(2024, 3, 31) for r in rows)
    assert all(len(r.sourceNotes) == 6 for r in rows)


def test_unknown_layout_does_not_silently_return_no_data():
    data = BytesIO()
    book = openpyxl.Workbook()
    book.active.append(["Unrecognized columns"])
    book.save(data)
    book.close()
    with Client() as client:
        for reader in (client.crr_hours, client.polr):
            with pytest.raises(ValueError, match="changed.xlsx/Sheet"):
                list(reader.read(data.getvalue(), filename="changed.xlsx"))
