# ruff: noqa: SIM117
"""Regression samples retain original cells, including formula-error tails."""

import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import pytest
from openpyxl import load_workbook

from tinyercot import Client

INPUTS = Path(__file__).resolve().parents[1] / "tools/inputs"
ROOT = INPUTS / "forecast-performance"
SAMPLES = json.loads((ROOT / "samples.json").read_text())


def sample_bytes(sample):
    with ZipFile(ROOT / "samples.zip") as archive:
        return archive.read(sample["member"])


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda s: s["title"])
def test_original_hourly_and_summary_cells(sample):
    data = sample_bytes(sample)
    with Client() as client:
        hourly = {
            (r.sourceSheet, r.sourceRow): r
            for r in client.load_forecast_performance.read(data)
        }
        summaries = {
            (r.sourceSheet, r.sourceRow, r.sourceColumn): r
            for r in client.load_forecast_performance.read_summaries(data)
        }
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    hours_checked = summaries_checked = 0
    for sheet in book:
        source = sheet.iter_rows(values_only=True)
        modern = next(source)[0] == "Date & Time"
        columns = (
            {
                1: "actual",
                2: "selected",
                3: "E",
                4: "E1",
                5: "E2",
                6: "E3",
                7: "M",
                8: "X",
                10: "sourceHour",
                11: "error",
                12: "under",
                13: "over",
            }
            if modern
            else {
                1: "actual",
                2: "selected",
                3: "A3",
                4: "A6",
                5: "E",
                6: "E1",
                7: "E2",
                8: "E3",
                9: "M",
                12: "sourceHour",
                13: "error",
                14: "under",
                15: "over",
                16: "frequencyUnder",
                17: "frequencyOver",
                18: "mape",
            }
        )
        chart_kind = "average_error"
        for line, cells in enumerate(source, 2):
            if any(c not in (None, "") for c in cells[: 14 if modern else 19]):
                row = hourly[sheet.title, line]
                for col, field in columns.items():
                    value = cells[col]
                    if isinstance(value, str) and value.startswith("#"):
                        assert getattr(row, field) is None
                        assert row.sourceMarkers[field] == value
                    else:
                        assert getattr(row, field) == (
                            None if value in (None, "") else Decimal(str(value))
                        )
                if isinstance(cells[0], str):
                    expected = datetime.strptime(  # noqa: DTZ007
                        cells[0], "%m/%d/%Y %H:%M" if modern else "%d%b%Y:%H:%M:%S"
                    )
                    assert row.timestamp == expected
                else:
                    assert row.timestamp is None
                assert row.hour == (
                    int(row.sourceHour)
                    if row.sourceHour is not None
                    and 1 <= row.sourceHour <= 24
                    and row.sourceHour == int(row.sourceHour)
                    else None
                )
                hours_checked += 1
            col = 15 if modern else 31 if sheet.title == "ERCOT" else 25
            if not modern and cells[col] in ("Hour", "HE"):
                chart_kind = "frequency" if cells[col] == "HE" else "average_error"
            columns2 = [(col, chart_kind)]
            if modern:
                columns2.append((19, "frequency"))
            elif cells[19] not in (None, ""):
                columns2.append((19, "monthly_mape"))
            for col, kind in columns2:
                if cells[col] in (None, "", "Hour", "HE"):
                    continue
                row = summaries[sheet.title, line, col + 1]
                assert row.kind == kind
                if kind == "monthly_mape":
                    assert row.hour is None
                    assert row.sourceBucket == "Monthly"
                    values = [("mape", cells[col])]
                else:
                    assert row.sourceBucket == str(cells[col])
                    assert row.hour == (
                        int(cells[col])
                        if isinstance(cells[col], (int, float))
                        else None
                    )
                    fields = (
                        ("frequencyUnder", "frequencyOver")
                        if kind == "frequency"
                        else ("under", "over")
                    )
                    values = list(zip(fields, cells[col + 1 : col + 3], strict=True))
                for field, value in values:
                    if isinstance(value, str) and value.startswith("#"):
                        assert getattr(row, field) is None
                        assert row.sourceMarkers[field] == value
                    else:
                        assert getattr(row, field) == (
                            None if value in (None, "") else Decimal(str(value))
                        )
                summaries_checked += 1
    book.close()
    assert hours_checked == len(hourly) > 0
    assert summaries_checked == len(summaries) > 0


def test_missing_actuals_and_calculation_only_rows_are_distinct():
    with Client() as client:
        service = client.load_forecast_performance
        january = next(s for s in SAMPLES if "January 2026" in s["title"])
        rows = list(
            service.read(sample_bytes(january), where=lambda r: r.region == "ERCOT")
        )
        assert rows[-1].timestamp == datetime.fromisoformat("2026-02-01T00:00:00")
        assert rows[-1].actual is None and rows[-1].selected is not None
        assert rows[-1].hour is None
        june = next(s for s in SAMPLES if "June 2025" in s["title"])
        tails = list(
            service.read(sample_bytes(june), where=lambda r: r.timestamp is None)
        )
        assert tails
        assert any(
            r.mape is None and r.sourceMarkers.get("mape") == "#DIV/0!" for r in tails
        )
        assert any(
            r.sourceMarkers.get("errorTimestamp") in ("0", "00:00:00") for r in tails
        )
        assert all(r.actual is None for r in tails)


def test_discovery_includes_mislabeled_metrics_and_attaches_source_file():
    sources = json.loads((ROOT / "sources.json").read_text())
    january = next(s for s in SAMPLES if "January 2026" in s["title"])
    with ZipFile(INPUTS / "peak-forecasts/indexes.zip") as archive:
        pages = {n: archive.read(n) for n in archive.namelist()}

    def handler(request):
        if request.url.path.startswith("/files/"):
            assert str(request.url) == january["url"]
            return httpx.Response(200, content=sample_bytes(january))
        year = request.url.path.rsplit("/", 1)[-1]
        return httpx.Response(
            200, content=pages[f"{year}.html" if year.isdigit() else "root.html"]
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as http:
        with Client(client=http) as client:
            service = client.load_forecast_performance
            files = service.files()
            assert {f.url for f in files} == {f["url"] for f in sources}
            service.files = lambda: [next(f for f in files if f.url == january["url"])]
            rows = list(
                service.rows(
                    where=lambda r: r.sourceFile is not None and r.region == "ERCOT"
                )
            )
            summaries = list(
                service.summaries(
                    where=lambda r: r.sourceFile is not None and r.kind == "frequency"
                )
            )
            assert rows and summaries
            assert rows[0].sourceFile.title == january["title"]
            assert summaries[0].hour == 1
