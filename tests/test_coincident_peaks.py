from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import openpyxl
import pytest

from tinyercot import Client, CoincidentPeakAllocation, MonthlyCoincidentPeak
from tinyercot._coincident_peaks import URL

FIXTURE = Path(__file__).resolve().parents[1] / "tools/inputs/four-cp/allocations.zip"
MONTHLY = FIXTURE.with_name("monthly.zip")


@pytest.fixture(scope="module")
def allocations():
    with Client() as client:
        return list(client.coincident_peaks.read_allocations(FIXTURE.read_bytes()))


def test_complete_captured_annual_tables(allocations):
    assert len(allocations) == 4345
    assert {r.year for r in allocations} == set(range(1996, 2021))
    assert Counter(r.kind for r in allocations) == {
        "allocation": 4212,
        "comparison": 124,
        "revision_detail": 9,
    }
    assert all(
        [p.month for p in r.peaks] == ([] if r.year == 1997 else [6, 7, 8, 9])
        for r in allocations
    )
    assert all(
        p.timestamp is None or p.timestamp.year == r.year
        for r in allocations
        for p in r.peaks
    )
    assert (
        CoincidentPeakAllocation.model_validate_json(allocations[0].model_dump_json())
        == allocations[0]
    )


def test_kw_values_and_historical_loss_adjustments(allocations):
    row = next(r for r in allocations if r.year == 1996 and r.entity == "Brazos")
    assert row.unit == "kW"
    assert row.averageLoad == Decimal("1061825.0")
    assert row.gsuLosses == Decimal("2062.9")
    assert row.lossAdjustmentFactor == Decimal("0.03519461844061518")
    assert row.peaks[0].load == Decimal("1107900.0")
    assert row.peaks[0].timestamp == datetime.fromisoformat("1996-06-19T17:00:00")
    assert row.peaks[0].timestamp.tzinfo is None
    assert row.duns is None
    assert row.adjustmentUnit == "factor"


def test_summary_below_monthly_tables_preserves_percent_units(allocations):
    row = next(r for r in allocations if r.year == 1997)
    assert row.entity == "Austin Electric Utility" and row.controlArea == "AEUX"
    assert row.peaks == []
    assert row.averageLoad == Decimal("1783073.3")
    assert row.adjustmentUnit == "percent"
    assert row.lossAdjustmentFactor == Decimal("1.87853")
    assert (
        row.sourceSection
        == "PRELIMINARY 1997 Four Coincident Peak Load Calculation as of 2/16/98"
    )


def test_revisions_are_retained_without_recomputing_shares(allocations):
    rows = [
        r
        for r in allocations
        if r.year == 2008
        and r.entity == "CENTERPOINT ENERGY HOUSTON ELECTRIC LLC (TDSP)"
    ]
    assert len(rows) == 2 and rows[0].sourceMember != rows[1].sourceMember
    assert rows[0].loadRatioShare == Decimal("0.25391237655230137")
    assert rows[1].oldLoadRatioShare == rows[0].loadRatioShare
    assert rows[1].loadRatioShareDifference == Decimal("0.000007873824533377594")
    assert rows[1].loadRatioShare == Decimal("0.25392025037683474")


def test_revision_sections_keep_year_and_missing_share(allocations):
    rows = [r for r in allocations if r.kind == "revision_detail"]
    assert all(r.year == 2019 for r in rows)
    assert rows[0].sourceSection == "ORIGINALLY FILED NUMBERS"
    revised = next(
        r
        for r in rows
        if r.sourceSection == "NEW AEP TEXAS CENTRAL LOAD PLUS Sharyland/McAllen"
    )
    assert revised.averageLoad == Decimal("5019.420838")
    assert revised.loadRatioShare is None
    assert revised.peaks[0].timestamp == datetime.fromisoformat("2019-06-19T17:00:00")


def test_modern_mw_and_string_identifiers(allocations):
    row = next(
        r
        for r in allocations
        if r.year == 2020 and r.entity == "AEP TEXAS CENTRAL COMPANY (TDSP)"
    )
    assert row.unit == "MW" and row.entityCode == "5"
    assert row.duns == "007924772"
    assert row.averageLoad == Decimal("4981.34814")
    assert row.peaks[-1].timestamp == datetime.fromisoformat("2020-09-01T14:30:00")
    assert row.gsuLosses is None


def test_missing_peak_clocks_are_not_inferred(allocations):
    row = next(r for r in allocations if r.year == 2001)
    assert all(p.timestamp is None for p in row.peaks)


def test_missing_first_month_does_not_drop_entity():
    with ZipFile(FIXTURE) as archive:
        name = next(n for n in archive.namelist() if n.endswith("Calculation2020.xlsx"))
        book = openpyxl.load_workbook(BytesIO(archive.read(name)))
    book.worksheets[0].cell(4, 4).value = None
    data = BytesIO()
    book.save(data)
    book.close()
    with Client() as client:
        row = next(client.coincident_peaks.read_allocations(data.getvalue()))
    assert row.entity == "AEP TEXAS CENTRAL COMPANY (TDSP)"
    assert row.peaks[0].load is None
    assert row.peaks[1].load == Decimal("5051.084372")


def test_anonymous_download_and_query_filters():
    requested = []

    def handler(request):
        assert "authorization" not in request.headers
        requested.append(str(request.url))
        return httpx.Response(200, content=FIXTURE.read_bytes())

    with (
        httpx.Client(transport=httpx.MockTransport(handler)) as http,
        Client(client=http) as client,
    ):
        rows = list(
            client.coincident_peaks.allocations(
                year_from=2020, year_to=2020, entity="AEP TEXAS CENTRAL COMPANY (TDSP)"
            )
        )
        assert len(rows) == 1 and rows[0].year == 2020
        with pytest.raises(ValueError, match="year_from"):
            list(client.coincident_peaks.allocations(year_from=2020, year_to=2019))
    assert requested == [URL]


@pytest.fixture(scope="module")
def monthly():
    with Client() as client:
        return list(client.coincident_peaks.read_monthly(MONTHLY.read_bytes()))


def test_monthly_tables_keep_every_publication(monthly):
    assert len(monthly) == 7633
    assert (
        len({(r.sourceMember, r.sourceSheet, r.sourceHeading) for r in monthly}) == 86
    )
    assert {r.timestamp.year for r in monthly} == set(range(1997, 2008))
    assert all(r.timestamp.tzinfo is None for r in monthly)
    assert (
        MonthlyCoincidentPeak.model_validate_json(monthly[0].model_dump_json())
        == monthly[0]
    )


def test_submitted_loads_missing_values_and_table_footnotes(monthly):
    rows = [
        r
        for r in monthly
        if r.timestamp.year == 1997 and r.entity == "Austin Electric Utility"
    ]
    june = rows[0]
    assert june.timestamp == datetime.fromisoformat("1997-06-30T17:00:00")
    assert june.loadType == "submitted" and june.unit == "kW"
    assert june.load == Decimal(1640586)
    assert june.gsuLossesKW == Decimal(2377)
    assert june.transmissionLossesKW is None
    assert june.controlAreaTotalKW == Decimal(1642963)
    assert june.demandReportRelativeDifference == Decimal("-0.000022519780888639415")
    assert june.sourceRow == 4
    assert any("metered load" in note for note in june.sourceNotes)
    assert not any("metered load" in note for note in rows[1].sourceNotes)
    assert june.settlementRun is None


def test_extended_submitted_report_values_and_row_note(monthly):
    row = next(r for r in monthly if r.timestamp.year == 1998 and r.sourceRowNote)
    assert row.entity == "Central Power & Light" and row.entityCode == "CPLC"
    assert row.controlArea == "CSWS"
    assert row.sisLoadReport == Decimal(5988)
    assert row.sisDifferenceMW == Decimal("-21.179399999999987")
    assert (
        row.sourceRowNote
        == "Rayburn Country (RCEC) was double counted in the D&E Report"
    )


def test_monthly_loss_adjustments_and_changed_peak_dates(monthly):
    row = next(
        r for r in monthly if r.timestamp.year == 2001 and r.entityCode == "ACPL"
    )
    assert row.loadType == "coincident_peak" and row.load == Decimal("3180561.6")
    assert row.lossAdjustmentPercent == Decimal("4.11072")
    assert row.loadAtDeliveryPointKW == Decimal("3054979.9")
    assert row.vamoLossPercent == Decimal("4.01376")
    assert row.loadResponsibilityKW == Decimal("3177599.5")
    assert {
        r.timestamp
        for r in monthly
        if r.timestamp.year == 2001 and r.timestamp.month == 8
    } == {
        datetime.fromisoformat("2001-08-15T16:30:00"),
        datetime.fromisoformat("2001-08-22T16:45:00"),
    }


def test_monthly_totals_and_unlabeled_source_rows(monthly):
    unnamed = [r for r in monthly if r.entity is None]
    assert len(unnamed) == 3
    assert unnamed[0].sourceRow == 59
    assert unnamed[0].load == Decimal("49274873.4")
    row = next(r for r in monthly if r.timestamp.year == 2001 and r.entity == "Total")
    assert row.vamoLossPercent is None
    assert row.loadResponsibilityKW == Decimal("50910818.1")
    row = next(r for r in monthly if r.timestamp.year == 2002 and r.entity == "Total")
    assert row.entityCode is None and row.duns is None
    assert row.load == Decimal("51805.597200000026")
    assert row.energyMWh == Decimal("12951.399300000006")


def test_settlement_stage_date_only_and_invalid_source_date(monthly):
    runs = [r.settlementRun for r in monthly if r.settlementRun]
    assert {r.kind for r in runs} == {None, "FINAL", "INTERIM", "TRUE-UP"}
    interim = next(r for r in runs if r.sourceDate == "1/22/04")
    assert interim.runDate == date(2004, 1, 22)
    assert interim.runTime is None
    typo = next(r for r in runs if r.sourceDate == "9/17/032")
    assert typo.runDate is None and typo.runTime == time(21, 48)
    unspecified = next(r for r in runs if r.sourceDate == "10/5/03")
    assert unspecified.channel == 5 and unspecified.kind is None


def test_monthly_filters_use_peak_date_despite_misnamed_sheet():
    with Client() as client:
        rows = list(
            client.coincident_peaks.read_monthly(
                MONTHLY.read_bytes(),
                date_from=date(2005, 8, 23),
                date_to=date(2005, 8, 23),
                entity="CENTERPOINT ENERGY HOUSTON ELECTRIC LLC (TDSP)",
            )
        )
    row = next(r for r in rows if r.sourceMember.endswith("2005-08CP.xls"))
    assert row.sourceSheet == "July"
    assert row.timestamp == datetime.fromisoformat("2005-08-23T16:30:00")
    assert row.load == Decimal("14802.5748") and row.unit == "MW"
    assert row.energyMWh == Decimal("3700.6437")
    assert row.sourceNotes == ["PRELIMINARY AUGUST 4CP"]
    assert row.settlementRun.kind == "FINAL"


def test_monthly_download_and_invalid_bounds():
    requested = []

    def handler(request):
        assert "authorization" not in request.headers
        requested.append(str(request.url))
        return httpx.Response(200, content=MONTHLY.read_bytes())

    with (
        httpx.Client(transport=httpx.MockTransport(handler)) as http,
        Client(client=http) as client,
    ):
        rows = list(
            client.coincident_peaks.monthly(
                date_from=date(1997, 6, 30),
                date_to=date(1997, 6, 30),
                entity="Austin Electric Utility",
            )
        )
        assert len(rows) == 1
        with pytest.raises(ValueError, match="date_from"):
            list(
                client.coincident_peaks.monthly(
                    date_from=date(2000, 1, 1), date_to=date(1999, 1, 1)
                )
            )
    assert requested == [URL]


def test_monthly_unknown_columns_report_source():
    book = openpyxl.Workbook()
    book.active.append(["Peak Interval Demand 06/13/2001 17:00"])
    book.active.append(["New header", "Unrecognized data"])
    data = BytesIO()
    book.save(data)
    book.close()
    with (
        Client() as client,
        pytest.raises(
            ValueError, match="saved.xlsx/Sheet: Missing monthly peak columns"
        ),
    ):
        list(
            client.coincident_peaks.read_monthly(data.getvalue(), filename="saved.xlsx")
        )
