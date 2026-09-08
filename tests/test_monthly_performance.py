"""Compare all published monthly table cells and their date/format semantics."""

import json
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import pytest
from openpyxl import load_workbook

from tinyercot import Client

INPUTS = Path(__file__).resolve().parents[1] / "tools/inputs"
ROOT = INPUTS / "monthly-performance"
SOURCES = json.loads((ROOT / "sources.json").read_text())


def original(source):
    with ZipFile(ROOT / "workbooks.zip") as archive:
        return archive.read(source["member"])


@pytest.mark.parametrize("source", SOURCES, ids=lambda f: f["title"])
def test_all_original_monthly_cells(source):
    data = original(source)
    with Client() as client:
        records = list(
            client.monthly_forecast_performance.read(data, filename=source["member"])
        )
    by_cell = {(r.sourceSheet, r.sourceRow, r.sourceColumn): r for r in records}
    assert len(by_cell) == len(records) == 222
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    specs = [("Forecast", 4, 1, range(5, 30), range(2, 7))]
    specs += [
        ("Backcast", header, month_col, range(header + 1, header + 13), [month_col + 1])
        for header in (5, 20)
        for month_col in (2, 5, 8)
    ]
    specs += [("Backcast", 5, 11, range(6, 31), [12])]
    checked = 0
    for sheet_name, header, month_col, lines, cols in specs:
        sheet = book[sheet_name]
        for line in lines:
            source_month = sheet.cell(line, month_col).value
            if isinstance(source_month, int):
                expected_month = date(
                    sheet.cell(header - 1, month_col).value, source_month, 1
                )
                expected_date = None
            else:
                expected_month = date(source_month.year, source_month.month, 1)
                expected_date = source_month.date()
            for col in cols:
                row = by_cell[sheet_name, line, col]
                cell = sheet.cell(line, col)
                assert row.sourceMember == source["member"]
                assert row.sourceMonthColumn == month_col
                assert row.month == expected_month and row.sourceDate == expected_date
                assert row.series == sheet.cell(header, col).value
                assert row.value == (
                    Decimal(str(cell.value)) if cell.value is not None else None
                )
                assert row.sourceNumberFormat == cell.number_format
                assert row.sourceMarker is None
                if cell.number_format == "0.00%" and row.value is not None:
                    assert row.percent == row.value * 100
                else:
                    assert row.percent is None
                if row.series in ("Goal", "Stretch"):
                    assert row.kind == "target"
                elif sheet_name == "Backcast":
                    assert row.kind == "backcast"
                else:
                    assert row.kind == "forecast"
                checked += 1
    book.close()
    assert checked == len(records)


def test_scales_dates_and_blank_future_months_remain_distinct():
    source = next(f for f in SOURCES if "August 2026" in f["title"])
    with Client() as client:
        rows = list(client.monthly_forecast_performance.read(original(source)))
    forecast = next(
        r for r in rows if r.month == date(2024, 8, 1) and r.series == "Day-Ahead RUC"
    )
    assert forecast.sourceDate == date(2024, 8, 24)
    assert forecast.value == Decimal("0.0141") and forecast.percent == Decimal("1.41")
    old = next(
        r for r in rows if r.month == date(2021, 1, 1) and r.series == "Backcast"
    )
    assert old.value == Decimal("1.366910678235441")
    assert old.percent is None and old.sourceDate is None
    future = next(r for r in rows if r.month == date(2026, 12, 1))
    assert future.value is None
    assert all(r.sourceRow != 101 for r in rows)


def test_live_index_shape_and_source_metadata():
    source = SOURCES[0]
    with ZipFile(INPUTS / "peak-forecasts/indexes.zip") as archive:
        pages = {n: archive.read(n) for n in archive.namelist()}

    def handler(request):
        if request.url.path.startswith("/files/"):
            assert str(request.url) == source["url"]
            return httpx.Response(200, content=original(source))
        year = request.url.path.rsplit("/", 1)[-1]
        return httpx.Response(
            200, content=pages[f"{year}.html" if year.isdigit() else "root.html"]
        )

    with (
        httpx.Client(transport=httpx.MockTransport(handler)) as http,
        Client(client=http) as client,
    ):
        reader = client.monthly_forecast_performance
        files = reader.files()
        assert {f.url for f in files} == {f["url"] for f in SOURCES}
        reader.files = lambda: [next(f for f in files if f.url == source["url"])]
        rows = list(
            reader.rows(
                where=lambda r: r.sourceFile is not None and r.kind == "forecast"
            )
        )
        assert len(rows) == 75
        assert rows[0].sourceFile.title == source["title"]
