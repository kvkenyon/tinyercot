from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import openpyxl
import pytest

from tinyercot import Client, LoadProfileDay
from tinyercot._load_profiles import ADJUSTMENTS_URL, INDEX_URL

FIXTURES = Path(__file__).resolve().parents[1] / "tools/inputs/load-profiles"


@pytest.mark.parametrize(
    "year,count", [(1997, 7), (2007, 4), (2012, 8), (2013, 13), (2026, 17)]
)
def test_captured_workbooks_cover_all_header_layouts(year, count):
    with Client() as client:
        rows = list(client.load_profiles.read((FIXTURES / f"{year}.zip").read_bytes()))
    assert len(rows) == count
    assert all(isinstance(row, LoadProfileDay) for row in rows)
    assert all(
        [i.interval for i in row.intervals] == list(range(1, 101)) for row in rows
    )
    assert all(
        row.profileType.startswith("BUS") and row.weatherZone in {"COAST", "EAST"}
        for row in rows
    )
    assert LoadProfileDay.model_validate_json(rows[0].model_dump_json()) == rows[0]


def test_dst_keeps_all_source_columns_and_blanks():
    with Client() as client:
        rows = list(client.load_profiles.read((FIXTURES / "2013.zip").read_bytes()))
    by_date = {row.operatingDay: row for row in rows}
    spring = by_date[date(2013, 3, 10)]
    autumn = by_date[date(2013, 11, 3)]
    assert sum(i.energyKWh is not None for i in spring.intervals) == 92
    assert all(i.energyKWh is None for i in spring.intervals[92:])
    assert sum(i.energyKWh is not None for i in autumn.intervals) == 100
    assert autumn.intervals[-1].energyKWh == Decimal("11.312")
    assert autumn.sourceAddTime == datetime.fromisoformat("2013-11-06T00:00:00")
    assert autumn.sourceAddTime.tzinfo is None


def test_auxiliary_profiles_remain_separate_and_saved_filename_is_retained():
    data = (FIXTURES / "2007.zip").read_bytes()
    with Client() as client:
        rows = list(client.load_profiles.read(data, profile="BUSHILF_COAST"))
        assert [(r.kind, r.intervals[0].energyKWh) for r in rows] == [
            ("backcast", Decimal("23.71")),
            ("auxiliary", Decimal("11.46")),
        ]
        assert rows[0].operatingDay == rows[1].operatingDay == date(2007, 1, 1)
        with ZipFile(BytesIO(data)) as archive:
            name = next(n for n in archive.namelist() if "Auxiliary" in n)
            saved = list(client.load_profiles.read(archive.read(name), filename=name))
        assert saved[0] == rows[1]


def test_xls_dates_and_optional_addtime():
    with Client() as client:
        old = next(client.load_profiles.read((FIXTURES / "1997.zip").read_bytes()))
        rows = list(client.load_profiles.read((FIXTURES / "2012.zip").read_bytes()))
    assert old.operatingDay == date(1997, 1, 3)
    assert old.intervals[0].energyKWh == Decimal("25.36")
    assert old.sourceAddTime is None
    assert rows[0].sourceAddTime is None
    q4 = next(r for r in rows if r.sourceSheet == "Q4 2012")
    assert q4.operatingDay == date(2012, 10, 1)
    assert q4.sourceAddTime == datetime.fromisoformat("2012-10-02T00:00:00")


def test_hurricane_factors_and_original_profiles_keep_distinct_units():
    data = (FIXTURES / "2008.zip").read_bytes()
    with Client() as client:
        rows = list(client.load_profiles.read(data))
        factors = list(client.load_profiles.read_adjustments(data))
        assert (
            list(
                client.load_profiles.read_adjustments(
                    (FIXTURES / "2013.zip").read_bytes()
                )
            )
            == []
        )
    original = [r for r in rows if r.kind == "original"]
    assert len(original) == 160
    assert all(len(r.intervals) == 96 for r in original)
    assert original[0].profile == "BUSHILF_COAST"
    assert original[0].intervals[0].energyKWh == Decimal("14.98")
    assert len(factors) == 16 * 96
    assert factors[0].weatherZone == "COAST"
    assert factors[0].operatingDay == date(2008, 9, 13)
    assert factors[0].factor == Decimal("0.3769911504424779")
    assert factors[-1].operatingDay == date(2008, 9, 28)
    assert factors[-1].interval == 96
    assert factors[-1].factor == Decimal("0.928598820058997")


def test_adjustment_query_downloads_directly_and_filters_date():
    requested = []

    def handler(request):
        requested.append(str(request.url))
        return httpx.Response(
            200,
            content=(
                (FIXTURES / "index.html").read_bytes()
                if str(request.url) == INDEX_URL
                else (FIXTURES / "2008.zip").read_bytes()
            ),
        )

    with (
        httpx.Client(transport=httpx.MockTransport(handler)) as http,
        Client(client=http) as client,
    ):
        rows = list(
            client.load_profiles.adjustments(
                date_from=date(2008, 9, 13), date_to=date(2008, 9, 13)
            )
        )
    assert requested == [ADJUSTMENTS_URL]
    assert len(rows) == 96 and all(r.operatingDay == date(2008, 9, 13) for r in rows)


def test_anonymous_discovery_and_bounded_query():
    requested = []

    def handler(request):
        assert "authorization" not in request.headers
        requested.append(str(request.url))
        return httpx.Response(
            200,
            content=(
                (FIXTURES / "index.html").read_bytes()
                if str(request.url) == INDEX_URL
                else (FIXTURES / "2026.zip").read_bytes()
            ),
        )

    with (
        httpx.Client(transport=httpx.MockTransport(handler)) as http,
        Client(client=http) as client,
    ):
        archives = client.load_profiles.archives()
        assert [a.year for a in archives] == list(range(1997, 2027))
        requested.clear()
        rows = list(
            client.load_profiles.rows(
                date_from=date(2026, 3, 8),
                date_to=date(2026, 3, 8),
                profile="BUSHIDG_COAST",
            )
        )
    assert requested == [INDEX_URL, archives[-1].url]
    assert len(rows) == 1 and rows[0].operatingDay == date(2026, 3, 8)
    assert rows[0].intervals[0].energyKWh == Decimal("12.831")


@pytest.mark.parametrize("method", ["rows", "read"])
def test_reversed_date_bounds(method):
    with Client() as client, pytest.raises(ValueError, match="date_from"):
        fn = getattr(client.load_profiles, method)
        list(
            fn(
                *([b""] if method == "read" else []),
                date_from=date(2026, 2, 1),
                date_to=date(2026, 1, 1),
            )
        )


@pytest.mark.parametrize("mutation", ["header", "extra"])
def test_unrecognized_columns_are_not_silently_dropped(mutation):
    with ZipFile(FIXTURES / "2026.zip") as archive:
        book = openpyxl.load_workbook(BytesIO(archive.read(archive.namelist()[0])))
    sheet = book.worksheets[0]
    if mutation == "header":
        sheet.cell(1, 3, "unexpected interval")
    else:
        sheet.cell(2, 104, 123)
    data = BytesIO()
    book.save(data)
    book.close()
    with Client() as client, pytest.raises(ValueError, match="columns|width"):
        list(client.load_profiles.read(data.getvalue()))
